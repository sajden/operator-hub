from __future__ import annotations

import json
import os
import sys
from urllib.error import URLError
from urllib.request import urlopen
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
except ModuleNotFoundError as exc:  # pragma: no cover - environment dependent
    raise SystemExit(
        "Missing dependency: openpyxl. Install with `pip install openpyxl` and retry."
    ) from exc

ROOT = Path(__file__).resolve().parents[1]
PROJECT_ID = os.environ.get("OPERATOR_HUB_PROJECT_ID", "parkpal")
GRAPH_PATH = ROOT / "projects" / PROJECT_ID / "data" / "graph.json"
OUTPUT_PATH = ROOT / "projects" / PROJECT_ID / "exports" / "Parkpal_Outreach.xlsx"
HUB_BASE_URL = os.environ.get("OPERATOR_HUB_URL", "http://127.0.0.1:8787")
HUB_PROJECTION_URL = f"{HUB_BASE_URL}/api/projects/{PROJECT_ID}/outreach-projection"
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.graph_projection import load_graph, project_graph

AUDIENCES_HEADERS = [
    "id",
    "project",
    "name",
    "parent_id",
    "relevance_why",
    "confidence_level",
    "status",
    "notes",
    "tags",
]

OUTREACH_HEADERS = [
    "id",
    "project",
    "audience_id",
    "channel",
    "community",
    "angle",
    "status",
    "notes",
]

PROBLEMS_HEADERS = [
    "id",
    "project",
    "parent_audience_id",
    "name",
    "confidence_level",
    "status",
    "notes",
    "tags",
]

CONTENT_HEADERS = [
    "id",
    "project",
    "audience_id",
    "outreach_id",
    "platform",
    "title",
    "body",
    "doc_link",
    "status",
    "sent_at",
    "posted_channel",
    "post_url",
]

FEEDBACK_HEADERS = [
    "id",
    "project",
    "audience_id",
    "content_id",
    "outreach_id",
    "source",
    "summary",
    "next_step",
    "status",
]

LINKS_HEADERS = ["id", "project", "type", "label", "url", "related_id"]


def populate_audiences(ws, rows: list[dict[str, Any]]) -> None:
    ws.append(AUDIENCES_HEADERS)
    for row in rows:
        ws.append([row.get(header, "") for header in AUDIENCES_HEADERS])


def populate_problems(ws, rows: list[dict[str, Any]]) -> None:
    ws.append(PROBLEMS_HEADERS)
    for row in rows:
        ws.append([row.get(header, "") for header in PROBLEMS_HEADERS])


def populate_outreach(ws, rows: list[dict[str, Any]]) -> None:
    ws.append(OUTREACH_HEADERS)
    for row in rows:
        ws.append([row.get(header, "") for header in OUTREACH_HEADERS])


def populate_content(ws, rows: list[dict[str, Any]]) -> None:
    ws.append(CONTENT_HEADERS)
    for row in rows:
        ws.append([row.get(header, "") for header in CONTENT_HEADERS])


def populate_feedback(ws, rows: list[dict[str, Any]]) -> None:
    ws.append(FEEDBACK_HEADERS)
    for row in rows:
        ws.append([row.get(header, "") for header in FEEDBACK_HEADERS])


def populate_links(ws, outreach_rows: list[dict[str, Any]]) -> None:
    ws.append(LINKS_HEADERS)
    for row in outreach_rows:
        url = row.get("url", "")
        if not url:
            continue
        ws.append(
            [
                f"link-{row.get('id', '')}",
                row.get("project", ""),
                "community",
                row.get("community", ""),
                url,
                row.get("id", ""),
            ]
        )


def load_projection() -> dict[str, list[dict[str, Any]]]:
    try:
        with urlopen(HUB_PROJECTION_URL, timeout=2) as response:
            return json.loads(response.read().decode("utf-8"))
    except (URLError, TimeoutError, OSError):
        graph = load_graph(GRAPH_PATH)
        return project_graph(graph)


def add_headers_only_sheet(workbook: Workbook, title: str, headers: list[str]) -> None:
    sheet = workbook.create_sheet(title=title)
    sheet.append(headers)


def main() -> None:
    try:
        projection = load_projection()

        workbook = Workbook()
        workbook.remove(workbook.active)

        audiences_sheet = workbook.create_sheet(title="Audiences")
        populate_audiences(audiences_sheet, projection["audiences"])

        problems_sheet = workbook.create_sheet(title="Problems")
        populate_problems(problems_sheet, projection["problems"])

        outreach_sheet = workbook.create_sheet(title="Outreach")
        populate_outreach(outreach_sheet, projection["outreach"])
        content_sheet = workbook.create_sheet(title="Content")
        populate_content(content_sheet, projection["content"])
        feedback_sheet = workbook.create_sheet(title="Feedback")
        populate_feedback(feedback_sheet, projection["feedback"])
        links_sheet = workbook.create_sheet(title="Links")
        populate_links(links_sheet, projection["outreach"])

        OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(OUTPUT_PATH)
        print(f"Exported workbook: {OUTPUT_PATH}")
    except Exception as exc:  # pragma: no cover - runtime validation path
        print(f"Export failed: {exc}", file=sys.stderr)
        raise SystemExit(1) from exc


if __name__ == "__main__":
    main()
